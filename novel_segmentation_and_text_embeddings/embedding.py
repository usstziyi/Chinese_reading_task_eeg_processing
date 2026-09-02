from transformers import BertTokenizer, AutoModelForMaskedLM
import openpyxl
import os
import numpy as np
import torch
import argparse


from transformers import AutoTokenizer, AutoModelForMaskedLM

# Load the pre-trained dictionary and tokenization method
tokenizer = BertTokenizer.from_pretrained(
    pretrained_model_name_or_path='bert-base-chinese',
)

model = AutoModelForMaskedLM.from_pretrained("bert-base-chinese")



Chinese_novels = {
    "LittlePrince": {
        'segmented_path':r'../data/segmented_novel_new/LittlePrince', 
        'run_num':7,
        'embedding_path':r'../data/embeddings/LittlePrince',
    },
    "GarnettDream": {
        'segmented_path':r'../data/segmented_novel_new/GarnettDream', 
        'run_num':18,
        'embedding_path':r'../data/embeddings/GarnettDream',
    }
}


parser = argparse.ArgumentParser(description='Parameters that can be changed in this experiment')
parser.add_argument('--novel_name',type=str,default='LittlePrince',
                    help='Novel key in Chinese_novels (e.g. LittlePrince, GarnettDream)')

args = parser.parse_args()

novel_cfg = Chinese_novels[args.novel_name]
args.segmented_path = novel_cfg['segmented_path']
args.run_num = novel_cfg['run_num']
args.embedding_path = novel_cfg['embedding_path']

os.makedirs(args.embedding_path, exist_ok=True)


model.eval()
for i in range(args.run_num):
    novel_path = args.segmented_path + '/segmented_Chinese_novel_run_' + str(i+1) + '.xlsx'

    wb = openpyxl.load_workbook(novel_path)
    wsheet = wb.active
    texts = []

    for j in range(2, wsheet.max_row + 1):
        texts.append((wsheet.cell(row=j, column=1)).value)


    embeddings = []
    for k in range(len(texts)):
        token = tokenizer.encode(texts[k], return_tensors='pt') # (1，seq_len)
        with torch.no_grad():
            # print(texts[k], token)
            embedding = model(token).logits # (1，seq_len，vocab_size)
            embedding = torch.mean(embedding, dim=1) # (1，vocab_size)
        embeddings.append(embedding.detach().numpy())


    embeddings = np.array(embeddings) # (num_texts, 1, vocab_size)

    embeddings = embeddings.reshape(embeddings.shape[0], embeddings.shape[2]) # (num_texts, vocab_size)


    np.save(args.embedding_path + '/text_embedding_run_' + str(i+1) + '.npy', embeddings)
    print(f'Run {i+1} embeddings saved')



